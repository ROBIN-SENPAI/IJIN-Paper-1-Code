# -*- coding: utf-8 -*-
"""Small toolkit to build a journal manuscript with python-docx: native Word equations (LaTeX subset
-> MathML -> OMML via Office's MML2OMML.XSL), IEEE numeric citations in order of first use, three-line
tables, numbered figures/tables/equations, line numbering and page numbers."""
import copy
import re

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm
from lxml import etree

XSL_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
_XSLT = None

GREEK = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "varepsilon": "ε", "zeta": "ζ",
         "eta": "η", "theta": "θ", "vartheta": "ϑ", "iota": "ι", "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν",
         "xi": "ξ", "pi": "π", "rho": "ρ", "sigma": "σ", "tau": "τ", "upsilon": "υ", "phi": "ϕ", "varphi": "φ",
         "chi": "χ", "psi": "ψ", "omega": "ω", "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ",
         "Pi": "Π", "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω", "ell": "ℓ"}
SYMS = {"leq": "≤", "geq": "≥", "le": "≤", "ge": "≥", "neq": "≠", "in": "∈", "notin": "∉", "cdot": "⋅",
        "times": "×", "to": "→", "rightarrow": "→", "leftarrow": "←", "infty": "∞", "sim": "∼", "approx": "≈",
        "ldots": "…", "cdots": "⋯", "mid": "∣", "subseteq": "⊆", "cup": "∪", "cap": "∩", "forall": "∀",
        "exists": "∃", "partial": "∂", "nabla": "∇", "pm": "±", "odot": "⊙", "oplus": "⊕", "Vert": "‖",
        "lVert": "‖", "rVert": "‖", "lfloor": "⌊", "rfloor": "⌋", "langle": "⟨", "rangle": "⟩", "wedge": "∧",
        "vee": "∨", "star": "⋆", "ast": "∗", "emptyset": "∅", "varnothing": "∅", "lceil": "⌈", "rceil": "⌉", "neg": "¬", "setminus": "∖", "prime": "′", "mapsto": "↦", "triangleq": "≜", "propto": "∝"}
FUNCS = {"log", "exp", "max", "min", "arg", "argmax", "argmin", "softmax", "sigmoid", "LeakyReLU", "ELU", "LN",
         "LSTM", "GAT", "MLP", "tanh", "sup", "inf", "median", "std", "mean", "Pr", "lim"}
BIG = {"sum": "∑", "prod": "∏"}
LIMIT_FUNCS = {"max", "min", "argmax", "argmin", "sup", "inf", "lim"}
BB = {"E": "𝔼", "R": "ℝ", "1": "𝟙", "N": "ℕ", "P": "ℙ"}


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class _TeX:
    def __init__(self, s):
        self.toks = re.findall(r"\\[a-zA-Z]+|\\.|[0-9]+(?:\.[0-9]+)?|[a-zA-Z]|\S", s)
        self.i = 0

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else None

    def next(self):
        t = self.peek()
        self.i += 1
        return t

    def group(self):
        t = self.next()
        if t == "{":
            out = self.expr("}")
            self.next()
            return "<mrow>" + out + "</mrow>"
        return self.atom(t)

    def raw_group(self):
        t = self.next()
        if t != "{":
            return t
        depth, buf = 1, []
        while True:
            t = self.next()
            if t == "{":
                depth += 1
            elif t == "}":
                depth -= 1
                if depth == 0:
                    break
            buf.append(t)
        return " ".join(buf).replace(" ", "") if all(len(b) == 1 for b in buf) else "".join(buf)

    def atom(self, t):
        if t is None:
            return ""
        if t.startswith("\\"):
            c = t[1:]
            if c == "frac":
                return f"<mfrac>{self.group()}{self.group()}</mfrac>"
            if c == "sqrt":
                return f"<msqrt>{self.group()}</msqrt>"
            if c in ("hat", "widehat"):
                return f'<mover accent="true">{self.group()}<mo>^</mo></mover>'
            if c in ("bar", "overline"):
                return f'<mover accent="true">{self.group()}<mo>¯</mo></mover>'
            if c in ("tilde", "widetilde"):
                return f'<mover accent="true">{self.group()}<mo>~</mo></mover>'
            if c == "dot":
                return f'<mover accent="true">{self.group()}<mo>˙</mo></mover>'
            if c == "mathcal":
                return f'<mi mathvariant="script">{_esc(self.raw_group())}</mi>'
            if c == "mathbb":
                g = self.raw_group()
                return f"<mi>{BB.get(g, g)}</mi>"
            if c in ("mathrm", "operatorname", "mathsf"):
                return f'<mi mathvariant="normal">{_esc(self.raw_group())}</mi>'
            if c == "mathbf":
                return f'<mi mathvariant="bold">{_esc(self.raw_group())}</mi>'
            if c == "boldsymbol":
                return f'<mstyle mathvariant="bold-italic">{self.group()}</mstyle>'
            if c == "text":
                return f"<mtext>{_esc(self.raw_group())}</mtext>"
            if c == "left":
                # \left d ... \right d2 -> a stretching delimiter around the enclosed expression (Word m:d)
                op = self._delim(self.next())
                inner = self.expr("\\right")
                self.next()
                cl = self._delim(self.next())
                return (f'<mfenced open="{_esc(op)}" close="{_esc(cl)}" separators="">'
                        f"<mrow>{inner}</mrow></mfenced>")
            if c in ("right", "big", "Big", "bigg", "Bigg", "bigl", "bigr", "Bigl", "Bigr"):
                d = self.next()
                if d in ("\\{", "\\}"):
                    d = d[1]
                if d == ".":
                    return ""
                if d.startswith("\\"):
                    d = SYMS.get(d[1:], d[1:])
                return f'<mo stretchy="true">{_esc(d)}</mo>'
            if c in (",", ";", ":", "quad", "qquad", " ", "!"):
                sp = {"quad": " ", "qquad": "  ", ";": " ", ":": " ", ",": " ", "!": "", " ": " "}[c]
                return f"<mtext>{sp}</mtext>" if sp else ""
            if c in ("{", "}"):
                return f"<mo>{c}</mo>"
            if c == "|":
                return "<mo>‖</mo>"
            if c in GREEK:
                return f"<mi>{GREEK[c]}</mi>"
            if c in SYMS:
                return f"<mo>{SYMS[c]}</mo>"
            if c in BIG:
                return f'<mo largeop="true" movablelimits="false">{BIG[c]}</mo>'
            if c in FUNCS:
                return f'<mi mathvariant="normal">{c}</mi>'
            return f"<mi>{_esc(c)}</mi>"
        if re.fullmatch(r"[0-9.]+", t):
            return f"<mn>{t}</mn>"
        if re.fullmatch(r"[a-zA-Z]", t):
            return f"<mi>{t}</mi>"
        if t == "-":
            return "<mo>−</mo>"
        if t == "'":
            return "<mo>′</mo>"
        if t == "{":
            out = self.expr("}")
            self.next()
            return "<mrow>" + out + "</mrow>"
        return f"<mo>{_esc(t)}</mo>"

    @staticmethod
    def _delim(d):
        if d is None or d == ".":
            return ""
        if d in ("\\{", "\\}"):
            return d[1]
        if d.startswith("\\"):
            return SYMS.get(d[1:], {"|": "‖"}.get(d[1:], d[1:]))
        return d

    def term(self, stop=None):
        t = self.next()
        base = self.atom(t)
        is_big = t.startswith("\\") and (t[1:] in BIG or t[1:] in LIMIT_FUNCS)
        sub = sup = None
        while self.peek() in ("_", "^"):
            op = self.next()
            if op == "_":
                sub = self.group()
            else:
                sup = self.group()
        if sub is not None and sup is not None:
            tag = "munderover" if is_big else "msubsup"
            base = f"<{tag}>{base}{sub}{sup}</{tag}>"
        elif sub is not None:
            tag = "munder" if is_big else "msub"
            base = f"<{tag}>{base}{sub}</{tag}>"
        elif sup is not None:
            tag = "mover" if is_big else "msup"
            base = f"<{tag}>{base}{sup}</{tag}>"
        if t.startswith("\\") and t[1:] in BIG:
            base = f"<mrow>{base}<mrow>{self.body(stop)}</mrow></mrow>"
        return base

    def body(self, stop):
        out, depth = [], 0
        while True:
            t = self.peek()
            if t is None or t == stop:
                break
            if depth == 0 and out and t in ("+", "-", "=", "<", ">", ",", r"\leq", r"\geq", r"\quad", r"\qquad", r"\le", r"\ge"):
                break
            if t in ("(", "[", r"\lceil", r"\lfloor", r"\langle"):
                depth += 1
            elif t in (")", "]", r"\right", r"\rceil", r"\rfloor", r"\rangle"):
                if depth == 0:
                    break
                depth -= 1
            out.append(self.term(stop))
        return "".join(out)

    def expr(self, stop=None):
        out = []
        while self.peek() is not None and self.peek() != stop:
            out.append(self.term(stop))
        return "".join(out)


def tex2omml(tex, size=None):
    global _XSLT
    if _XSLT is None:
        _XSLT = etree.XSLT(etree.parse(XSL_PATH))
    mml = '<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow>' + _TeX(tex).expr() + "</mrow></math>"
    res = _XSLT(etree.fromstring(mml.encode("utf-8")))
    root = copy.deepcopy(res.getroot())
    if size:
        M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        Wn = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        for r in root.iter("{%s}r" % M):
            rpr = etree.SubElement(r, "{%s}rPr" % Wn)
            r.remove(rpr)
            t = r.find("{%s}t" % M)
            idx = list(r).index(t) if t is not None else len(r)
            r.insert(idx, rpr)
            for tag in ("sz", "szCs"):
                e = etree.SubElement(rpr, "{%s}%s" % (Wn, tag))
                e.set("{%s}val" % Wn, str(int(round(size * 2))))
    return root


# ----------------------------------------------------------------------------------------- document
class Manuscript:
    def __init__(self, refs, font="Times New Roman", size=10, spacing=2.0, line_numbers=True, prelabels=None,
                 prefix="", supp=None, main_labels=None, parent=None, secmap=None, heading_names=None, uniform=None,
                 wide_landscape=False, cite_style="ieee", align="justify", page="A4", margins_cm=2.5):
        """secmap renumbers headings and 'Section N.M' references ({old top-level number: new prefix});
        uniform forces one font size for all text (captions, tables, references); wide_landscape places tables with
        at least that many columns (True = 8) on landscape pages."""
        self.doc = Document()
        self.cite_style, self.default_align = cite_style, align
        self.secmap, self.heading_names = secmap or {}, heading_names or {}
        self.uniform, self.wide_landscape = uniform, wide_landscape
        self.cap_size = uniform or 9
        self.cap_spacing = spacing if uniform else 1.15
        self.fig_files = []
        self.refs = refs
        self.cite_order = []
        self.fig_n = self.tab_n = self.eq_n = 0
        self._after_eq = False      # a paragraph that continues the sentence of an equation is not indented
        self.labels = dict(prelabels or {})
        self.prefix, self.supp, self.parent = prefix, supp, parent
        self.main_labels = main_labels
        self.font, self.size, self.spacing = font, size, spacing
        st = self.doc.styles["Normal"]
        st.font.name = font
        st.font.size = Pt(size)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        st.element.rPr.rFonts.set(qn("w:cs"), font)
        pf = st.paragraph_format
        pf.line_spacing = spacing
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        sec = self.doc.sections[0]
        if page == "Letter":
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        else:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(margins_cm)
        sec.top_margin = sec.bottom_margin = Cm(margins_cm)
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1")
            ln.set(qn("w:restart"), "continuous")
            ln.set(qn("w:distance"), "283")
            sec._sectPr.append(ln)
        self._footer_page_number(sec)
        for lvl, (sz, bold, italic) in {1: (11, True, False), 2: (10, True, True), 3: (10, False, True)}.items():
            hs = self.doc.styles[f"Heading {lvl}"]
            hs.font.name = font
            hs.font.size = Pt(uniform or sz)
            hs.font.bold = bold
            hs.font.italic = italic
            hs.font.color.rgb = RGBColor(0, 0, 0)
            rpr = hs.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), font)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            hs.paragraph_format.space_before = Pt(12 if lvl == 1 else 6)
            hs.paragraph_format.space_after = Pt(6 if lvl == 1 else 3)
            hs.paragraph_format.line_spacing = spacing
            hs.paragraph_format.keep_with_next = True

    def _footer_page_number(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if typ:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), typ)
                r._r.append(e)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
                r._r.append(e)
        r.font.size = Pt(self.uniform or 9)

    # -------------------------------------------------------------- section renumbering
    def _map_num(self, num):
        parts = num.split(".")
        if not self.secmap or parts[0] not in self.secmap:
            return num
        return ".".join([self.secmap[parts[0]]] + parts[1:])

    def _map_secrefs(self, text):
        if not self.secmap:
            return text

        def fix(m):
            return re.sub(r"\d+(?:\.\d+)*", lambda n: self._map_num(n.group(0)), m.group(0))
        return re.sub(r"\bSections? \d+(?:\.\d+)*(?:(?:, | and | or |–)\d+(?:\.\d+)*)*", fix, text)

    # -------------------------------------------------------------- citations and inline markup
    def _routed(self, kind, label):
        return self.supp is not None and label is not None and f"{kind}:{label}" not in self.main_labels

    def cite(self, keys):
        if self.parent is not None:
            return self.parent.cite(keys)
        nums = []
        for k in keys:
            k = k.strip().lstrip("@")
            if k not in self.refs:
                raise KeyError(f"unknown reference {k}")
            if k not in self.cite_order:
                self.cite_order.append(k)
            nums.append(self.cite_order.index(k) + 1)
        nums = sorted(set(nums))
        groups, run = [], [nums[0]]
        for n in nums[1:]:
            if n == run[-1] + 1:
                run.append(n)
            else:
                groups.append(run)
                run = [n]
        groups.append(run)
        parts = []
        for g in groups:
            if len(g) >= 3:
                parts.append(f"{g[0]}–{g[-1]}")
            else:
                parts.extend(str(n) for n in g)
        if self.cite_style == "elsevier":
            return "[" + ",".join(parts) + "]"
        return ", ".join(f"[{x}]" if "–" not in x else "[" + x.replace("–", "]–[") + "]" for x in parts)

    def _resolve(self, text):
        text = re.sub(r"\[@([^\]]+)\]", lambda m: self.cite(m.group(1).split(";")), text)
        text = re.sub(r"\{(fig|tab|eq|sec|pan):([a-zA-Z0-9_]+)\}", lambda m: str(self.labels.get(m.group(1) + ":" + m.group(2), "??")), text)
        return self._map_secrefs(text)

    def _runs(self, p, text, size=None, italic=False, bold=False, color=None):
        text = self._resolve(text)
        for seg in re.split(r"(\$[^$]+\$)", text):
            if not seg:
                continue
            if seg.startswith("$") and seg.endswith("$") and len(seg) > 2:
                p._p.append(tex2omml(seg[1:-1], size=size))
                continue
            for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\^\{[^}]+\}|_\{[^}]+\})", seg):
                if not part:
                    continue
                b, it, sup, sub = bold, italic, False, False
                if part.startswith("**"):
                    part, b = part[2:-2], True
                elif part.startswith("*") and part.endswith("*") and len(part) > 1:
                    part, it = part[1:-1], not italic
                elif part.startswith("^{"):
                    part, sup = part[2:-1], True
                elif part.startswith("_{"):
                    part, sub = part[2:-1], True
                r = p.add_run(part)
                r.bold = True if b else None
                r.italic = True if it else None
                if sup:
                    r.font.superscript = True
                if sub:
                    r.font.subscript = True
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = RGBColor.from_string(color)
        return p

    # -------------------------------------------------------------- blocks
    def para(self, text, align=None, indent=True, size=None, italic=False, bold=False, space_after=0, keep=False,
             spacing=None, color=None):
        if self.uniform and size:
            size = self.uniform
        align = align or self.default_align
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        if self._after_eq and re.match(r"[a-z]", text.lstrip()):
            indent = False
        self._after_eq = False
        if indent:
            p.paragraph_format.first_line_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(space_after)
        if spacing:
            p.paragraph_format.line_spacing = spacing
        if keep:
            p.paragraph_format.keep_with_next = True
        return self._runs(p, text, size=size, italic=italic, bold=bold, color=color)

    def heading(self, text, level=1, label=None):
        self._after_eq = False
        m = re.match(r"(\d+(?:\.\d+)*)\.? (.*)", text)
        if self.secmap and m:
            num = self._map_num(m.group(1))
            title = self.heading_names.get(m.group(2), m.group(2))
            text, level = f"{num}. {title}", min(3, num.count(".") + 1)
        h = self.doc.add_heading(level=level)
        self._runs(h, text)
        if label:
            self.labels["sec:" + label] = text.split(" ")[0].rstrip(".")
        return h

    def bullets(self, items, numbered=False):
        for i, t in enumerate(items):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.default_align == 'left' else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            mark = f"({i + 1})\u00a0" if numbered else "•\u00a0"
            r = p.add_run(mark)
            self._runs(p, t)

    def equation(self, tex, label=None):
        self.eq_n += 1
        if label:
            self.labels["eq:" + label] = f"({self.eq_n})"
        # borderless 1 x 3 table: display-mode equation centred, number flush right
        t = self.doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = self.text_width
        widths = [int(w * 0.08), int(w * 0.84), int(w * 0.08)]
        for j, cell in enumerate(t.rows[0].cells):
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            if j == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para = OxmlElement("m:oMathPara")
                para.append(tex2omml(tex))
                p._p.append(para)
            elif j == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.add_run(f"({self.eq_n})")
            cell.vertical_alignment = 1
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "nil")
            borders.append(e)
        tblPr.append(borders)
        self._after_eq = True
        return t

    def label_next(self, kind, label):
        if self._routed(kind, label):
            self.supp.label_next(kind, label)
            self.labels[f"{kind}:{label}"] = self.supp.labels[f"{kind}:{label}"]
            return
        n = {"fig": self.fig_n, "tab": self.tab_n}[kind] + 1
        self.labels[f"{kind}:{label}"] = f"{self.prefix}{n}"

    def figure(self, path, caption, width_in=6.3, label=None):
        if self._routed("fig", label):
            self.supp.labels.update(self.labels)
            self.supp.figure(path, caption, width_in, label)
            self.labels["fig:" + label] = self.supp.labels["fig:" + label]
            return
        self.fig_n += 1
        self._after_eq = False
        if label:
            self.labels["fig:" + label] = f"{self.prefix}{self.fig_n}"
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(path, width=Inches(width_in))
        self.fig_files.append((self.fig_n, path, width_in))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.default_align == 'left' else WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.space_after = Pt(8)
        c.paragraph_format.keep_together = True        # a caption is never split across pages
        c.paragraph_format.line_spacing = self.cap_spacing
        r = c.add_run(f"Fig. {self.prefix}{self.fig_n}. ")
        r.bold = True
        r.font.size = Pt(self.cap_size)
        self._runs(c, caption, size=self.cap_size)

    def table(self, header, rows, caption, col_widths=None, label=None, size=8, align=None, notes=None, header_rows=1):
        if self._routed("tab", label):
            self.supp.labels.update(self.labels)
            t = self.supp.table(header, rows, caption, col_widths, label, size, align, notes, header_rows)
            self.labels["tab:" + label] = self.supp.labels["tab:" + label]
            return t
        self.tab_n += 1
        self._after_eq = False
        if label:
            self.labels["tab:" + label] = f"{self.prefix}{self.tab_n}"
        allrows = ([header] if header_rows == 1 else list(header)) + rows
        ncol = len(allrows[0])
        land = bool(self.wide_landscape) and ncol >= (8 if self.wide_landscape is True else self.wide_landscape)
        tw = self._orient(True) if land else self.text_width
        if self.uniform:
            size = self.uniform
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.default_align == 'left' else WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_before = Pt(8)
        c.paragraph_format.space_after = Pt(3)
        c.paragraph_format.line_spacing = self.cap_spacing
        r = c.add_run(f"Table {self.prefix}{self.tab_n}. ")
        r.bold = True
        r.font.size = Pt(self.cap_size)
        self._runs(c, caption, size=self.cap_size)
        t = self.doc.add_table(rows=len(allrows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if col_widths is None:
            col_widths = [tw / ncol] * ncol
        else:
            tot = sum(col_widths)
            col_widths = [int(tw * w / tot) for w in col_widths]
        # keep short tables on one page (and their last row with the note); long tables may break between rows
        keep_all = len(allrows) <= 24
        for i, row in enumerate(allrows):
            trPr = t.rows[i]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i < header_rows:
                trPr.append(OxmlElement("w:tblHeader"))      # header rows repeat on every page
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = col_widths[j]
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                if i < header_rows or (keep_all and (i < len(allrows) - 1 or notes)):
                    p.paragraph_format.keep_with_next = True
                a = (align[j] if align else ("left" if j == 0 else "center"))
                p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                               "right": WD_ALIGN_PARAGRAPH.RIGHT}[a]
                sval = str(val)
                if re.fullmatch(r"\$[^$]+\$", sval.strip()):
                    sval = "​" + sval      # a cell holding only an equation would be centred as display math
                self._runs(p, sval, size=size, bold=(i < header_rows))
                tcPr = cell._tc.get_or_add_tcPr()
                mar = OxmlElement("w:tcMar")
                for side in ("top", "bottom"):
                    e = OxmlElement(f"w:{side}")
                    e.set(qn("w:w"), "28")
                    e.set(qn("w:type"), "dxa")
                    mar.append(e)
                tcPr.append(mar)
        if header_rows == 2:
            self._merge_group_header(t, allrows, ncol)
        self._three_line(t, header_rows)
        if notes:
            n = self.doc.add_paragraph()
            n.paragraph_format.line_spacing = self.cap_spacing if self.uniform else 1.0
            n.paragraph_format.space_after = Pt(8)
            self._runs(n, notes, size=self.uniform or 7.5)
        elif not land:
            sp = self.doc.add_paragraph()
            sp.paragraph_format.line_spacing = 1.0
        if land:
            self._orient(False)
        return t

    def _orient(self, landscape):
        """Start a new page-section in landscape or portrait and return its text width. Line numbering and the
        footer are inherited from the previous section."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        a, b = Cm(29.7), Cm(21.0)
        sec.page_width, sec.page_height = (a, b) if landscape else (b, a)
        return sec.page_width - sec.left_margin - sec.right_margin

    @staticmethod
    def _merge_group_header(t, allrows, ncol):
        """Two header rows: a group label followed by empty cells spans the sub-columns below it (with a thin rule
        under it); a label with nothing below it spans both header rows."""
        top, sub = [str(v) for v in allrows[0]], [str(v) for v in allrows[1]]
        j = 0
        while j < ncol:
            k = j
            while k + 1 < ncol and top[k + 1] == "" and sub[k + 1] != "":
                k += 1
            if k > j and top[j] != "":
                mc = t.cell(0, j).merge(t.cell(0, k))
                mc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                tcPr = mc._tc.get_or_add_tcPr()
                b = OxmlElement("w:tcBorders")
                e = OxmlElement("w:bottom")
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), "000000")
                b.append(e)
                tcPr.append(b)
            j = k + 1
        for j in range(ncol):
            if top[j] != "" and sub[j] == "":
                t.cell(0, j).merge(t.cell(1, j))

    def _three_line(self, t, header_rows):
        tbl = t._tbl
        tblPr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), "000000")
            borders.append(e)
        for edge in ("left", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "nil")
            borders.append(e)
        tblPr.append(borders)
        for cell in t.rows[header_rows - 1].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            e = OxmlElement("w:bottom")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "6")
            e.set(qn("w:color"), "000000")
            b.append(e)
            tcPr.append(b)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def references(self, title="References"):
        h = self.doc.add_heading(level=1)
        h.add_run(title)
        for i, k in enumerate(self.cite_order):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if self.default_align == 'left' else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.line_spacing = self.spacing
            r = p.add_run(f"[{i + 1}]\t")
            r.font.size = Pt(self.cap_size)
            r2 = p.add_run(self.refs[k])
            r2.font.size = Pt(self.cap_size)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))

    def save(self, path, title=None):
        cp = self.doc.core_properties  # python-docx fills 'python-docx' as author; keep the file free of identities
        cp.author = cp.last_modified_by = ""
        cp.comments = ""
        if title:
            cp.title = title
        self.doc.save(path)
