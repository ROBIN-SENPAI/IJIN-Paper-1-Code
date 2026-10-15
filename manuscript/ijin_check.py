# -*- coding: utf-8 -*-
"""Programmatic verification of IJIN_formatted.docx against the IJIN formatting instructions."""
import re
import zipfile

from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _blocks(doc):
    for el in doc.element.body.iterchildren():
        if el.tag == W + "p":
            yield "p", "".join(t.text or "" for t in el.iter(W + "t")), el
        elif el.tag == W + "tbl":
            yield "t", " ".join("".join(t.text or "" for t in p.iter(W + "t")) for p in el.iter(W + "p")), el


def _expand(txt):
    out = []
    for g in re.finditer(r"\[(\d+(?:[,–]\d+)*)\]", txt):
        for part in g.group(1).split(","):
            if "–" in part:
                a, b = part.split("–")
                out.extend(range(int(a), int(b) + 1))
            else:
                out.append(int(part))
    return out


def verify(path, pdf_path=None):
    C = {}
    doc = Document(path)
    body = doc.element.body
    C["reopen"] = True
    C["justified_paragraphs"] = sum(1 for j in body.iter(W + "jc") if j.get(W + "val") in ("both", "distribute"))
    settings = zipfile.ZipFile(path).read("word/settings.xml").decode("utf-8")
    C["auto_hyphenation"] = "autoHyphenation" in settings and 'w:autoHyphenation w:val="0"' not in settings
    secs = list(body.iter(W + "sectPr"))
    C["multi_column_sections"] = sum(1 for s in secs if s.find(W + "cols") is not None and
                                     int(s.find(W + "cols").get(W + "num", "1")) > 1)
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    C["headings"] = heads
    num = [t for s, t in heads if re.match(r"^\d", t)]
    C["numbered_headings"] = num
    C["heading_format_ok"] = all(re.match(r"^\d+\. \S", t) if s == "Heading 1" else re.match(r"^\d+(\.\d+)+ \S", t)
                                 for s, t in heads if re.match(r"^\d", t))
    C["display_equations"] = sum(1 for _ in body.iter(M + "oMathPara"))
    C["math_objects"] = sum(1 for _ in body.iter(M + "oMath"))
    C["images"] = len(doc.inline_shapes)
    blocks = list(_blocks(doc))
    ref_i = next(k for k, (kd, t, _) in enumerate(blocks) if kd == "p" and t.strip() == "References")
    figcap, tabcap, fig_first, tab_first, cites = [], [], {}, {}, []
    for k, (kd, t, el) in enumerate(blocks[:ref_i]):
        mf = re.match(r"Fig\. (\d+)\. ", t) if kd == "p" else None
        mt = re.match(r"Table (\d+)\. ", t) if kd == "p" else None
        if mf:
            figcap.append(int(mf.group(1))); continue
        if mt:
            tabcap.append(int(mt.group(1))); continue
        for g in re.finditer(r"Figs?\. (\d+)", t):
            fig_first.setdefault(int(g.group(1)), k)
        for g in re.finditer(r"\bTables? (\d+)", t):
            tab_first.setdefault(int(g.group(1)), k)
        cites.extend(_expand(t))
    C["figure_captions"], C["table_captions"] = figcap, tabcap
    C["figures_sequential"] = figcap == list(range(1, len(figcap) + 1))
    C["tables_sequential"] = tabcap == list(range(1, len(tabcap) + 1))
    C["figures_not_cited"] = [n for n in figcap if n not in fig_first]
    C["tables_not_cited"] = [n for n in tabcap if n not in tab_first]
    first = []
    for n in cites:
        if n not in first:
            first.append(n)
    reflist = [t for kd, t, _ in blocks[ref_i + 1:] if kd == "p" and re.match(r"\[\d+\]", t)]
    C["references_listed"] = len(reflist)
    C["citations_in_order"] = first == list(range(1, len(first) + 1))
    C["cited_not_listed"] = [n for n in first if n > len(reflist)]
    C["listed_not_cited"] = [n for n in range(1, len(reflist) + 1) if n not in first]
    full = "\n".join(t for _, t, _ in blocks)
    C["tags"] = {t: full.count(t) for t in ("[AUTHOR INFORMATION REQUIRED", "[METHODOLOGICAL DETAIL REQUIRED]",
                                             "[CAPTION OR SOURCE REQUIRED]") if t in full}
    C["leftover_placeholders"] = sorted(set(re.findall(r"\?\?|\{(?:fig|tab|eq|sec|pan):[a-z0-9_]+\}", full)))
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        C["comments_part"] = any("comments" in n for n in z.namelist())
    C["tracked_changes"] = len(re.findall(r"<w:(ins|del) ", xml))
    C["core_author"] = doc.core_properties.author
    if pdf_path:
        from pypdf import PdfReader
        rd = PdfReader(pdf_path)
        C["pdf_pages"] = len(rd.pages)
        low = []
        for k, pg in enumerate(rd.pages, 1):
            txt = re.sub(r"\s+", " ", pg.extract_text() or "").strip()
            if len([w for w in txt.split() if not w.isdigit()]) < 15 and not getattr(pg, "images", []):
                low.append(k)
        C["pdf_blank_pages"] = low
    return C
